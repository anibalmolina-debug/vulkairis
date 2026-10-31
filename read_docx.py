#!/usr/bin/env python3
import docx
doc = docx.Document(r"C:\Users\Molina Nuñez\Downloads\Vulkairis_RunClient_Fix_Guide.docx")
for para in doc.paragraphs:
    print(f"[{para.style.name}] {para.text}")
for table in doc.tables:
    for row in table.rows:
        print(" | ".join([cell.text for cell in row.cells]))